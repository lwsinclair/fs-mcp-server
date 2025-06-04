import io
import os
import logging
import warnings
from pathlib import Path
from typing import Union, Optional

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

# 抑制PDF处理警告
warnings.filterwarnings('ignore', category=UserWarning, module='pdfminer')
warnings.filterwarnings('ignore', message='.*CropBox missing.*')

# 设置pdfminer日志级别
logging.getLogger("pdfminer").setLevel(logging.ERROR)
logging.getLogger("pdfplumber").setLevel(logging.ERROR)

class FileConverter:
    """文档格式转换器"""

    @staticmethod
    def docx_to_markdown(file_path: Union[str, Path]) -> str:
        """
        将DOCX文件转换为Markdown格式

        Args:
            file_path: DOCX文件路径

        Returns:
            str: Markdown格式的文本

        Raises:
            ImportError: 如果python-docx库未安装
            Exception: 文件读取错误
        """
        if Document is None:
            raise ImportError("需要安装python-docx库: pip install python-docx")

        try:
            doc = Document(file_path)
            markdown_content = []

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    # 简单的格式转换
                    if paragraph.style.name.startswith("Heading"):
                        level = (
                            int(paragraph.style.name.split()[-1])
                            if paragraph.style.name.split()[-1].isdigit()
                            else 1
                        )
                        markdown_content.append(f"{'#' * level} {text}")
                    else:
                        markdown_content.append(text)
                    markdown_content.append("")  # 添加空行

            # 处理表格
            for table in doc.tables:
                markdown_content.append("")  # 表格前空行
                for i, row in enumerate(table.rows):
                    cells = [cell.text.strip() for cell in row.cells]
                    markdown_content.append("| " + " | ".join(cells) + " |")
                    if i == 0:  # 添加表头分隔线
                        markdown_content.append(
                            "| " + " | ".join(["---"] * len(cells)) + " |"
                        )
                markdown_content.append("")  # 表格后空行

            return "\n".join(markdown_content)

        except Exception as e:
            raise Exception(f"DOCX文件转换失败: {e}")

    @staticmethod
    def xlsx_to_markdown(file_path: Union[str, Path]) -> str:
        """
        将XLSX文件转换为Markdown格式

        Args:
            file_path: XLSX文件路径

        Returns:
            str: Markdown格式的文本

        Raises:
            ImportError: 如果pandas库未安装
            Exception: 文件读取错误
        """
        if pd is None:
            raise ImportError("需要安装pandas库: pip install pandas openpyxl")

        try:
            # 读取所有工作表
            excel_file = pd.ExcelFile(file_path)
            markdown_content = []

            for sheet_name in excel_file.sheet_names:
                markdown_content.append(f"# {sheet_name}")
                markdown_content.append("")

                df = pd.read_excel(file_path, sheet_name=sheet_name)

                # 将DataFrame转换为Markdown表格
                if not df.empty:
                    # 创建表头
                    headers = df.columns.tolist()
                    markdown_content.append(
                        "| " + " | ".join(str(h) for h in headers) + " |"
                    )
                    markdown_content.append(
                        "| " + " | ".join(["---"] * len(headers)) + " |"
                    )

                    # 添加数据行
                    for _, row in df.iterrows():
                        cells = [str(cell) if pd.notna(cell) else "" for cell in row]
                        markdown_content.append("| " + " | ".join(cells) + " |")

                markdown_content.append("")  # 工作表之间的空行

            return "\n".join(markdown_content)

        except Exception as e:
            raise Exception(f"XLSX文件转换失败: {e}")

    @staticmethod
    def pdf_to_markdown(
        file_path: Union[str, Path], use_pdfplumber: bool = True
    ) -> str:
        """
        将PDF文件转换为Markdown格式

        Args:
            file_path: PDF文件路径
            use_pdfplumber: 是否使用pdfplumber库（更好的文本提取）

        Returns:
            str: Markdown格式的文本

        Raises:
            ImportError: 如果所需库未安装
            Exception: 文件读取错误
        """
        if use_pdfplumber and pdfplumber is not None:
            return FileConverter._pdf_to_markdown_pdfplumber(file_path)
        elif PyPDF2 is not None:
            return FileConverter._pdf_to_markdown_pypdf2(file_path)
        else:
            raise ImportError(
                "需要安装PDF处理库: pip install pdfplumber 或 pip install PyPDF2"
            )

    @staticmethod
    def _pdf_to_markdown_pdfplumber(file_path: Union[str, Path]) -> str:
        """使用pdfplumber提取PDF文本，专注文本内容忽略媒体信息"""
        try:
            markdown_content = []

            # 抑制特定的警告
            with warnings.catch_warnings():
                warnings.filterwarnings('ignore', message='.*CropBox missing.*')
                warnings.filterwarnings('ignore', category=UserWarning, module='pdfminer')
                
                with pdfplumber.open(file_path) as pdf:
                    for i, page in enumerate(pdf.pages):
                        if i > 0:  # 添加页面分隔
                            markdown_content.append(f"\n---\n# 第 {i + 1} 页\n")

                        # 配置文本提取参数，专注文本内容
                        text = page.extract_text(
                            x_tolerance=3,
                            y_tolerance=3,
                            layout=False,  # 不保持布局，简化输出
                            x_density=7.25,
                            y_density=13
                        )
                        
                        if text:
                            # 清理和格式化文本
                            text = text.strip()
                            # 移除多余的空行
                            text = '\n'.join(line.strip() for line in text.split('\n') if line.strip())
                            
                            # 简单的段落处理
                            paragraphs = text.split('\n\n')
                            for paragraph in paragraphs:
                                paragraph = paragraph.strip()
                                if paragraph and len(paragraph) > 3:  # 过滤太短的内容
                                    markdown_content.append(paragraph)
                                    markdown_content.append("")

                        # 提取表格（如果需要）
                        try:
                            tables = page.extract_tables()
                            for table in tables:
                                if table and len(table) > 0:
                                    markdown_content.append("")
                                    for j, row in enumerate(table):
                                        if row and any(cell for cell in row if cell):  # 确保行不为空
                                            cells = [str(cell).strip() if cell else "" for cell in row]
                                            if any(cells):  # 确保不是空行
                                                markdown_content.append(
                                                    "| " + " | ".join(cells) + " |"
                                                )
                                                if j == 0:  # 表头分隔线
                                                    markdown_content.append(
                                                        "| "
                                                        + " | ".join(["---"] * len(cells))
                                                        + " |"
                                                    )
                                    markdown_content.append("")
                        except Exception:
                            # 如果表格提取失败，忽略表格但继续处理
                            pass

            # 清理最终内容
            final_content = '\n'.join(markdown_content).strip()
            # 移除多余的连续空行
            while '\n\n\n' in final_content:
                final_content = final_content.replace('\n\n\n', '\n\n')
            
            return final_content

        except Exception as e:
            raise Exception(f"PDF文件转换失败 (pdfplumber): {e}")

    @staticmethod
    def _pdf_to_markdown_pypdf2(file_path: Union[str, Path]) -> str:
        """使用PyPDF2提取PDF文本"""
        try:
            markdown_content = []

            with open(file_path, "rb") as file:
                pdf_reader = PyPDF2.PdfReader(file)

                for i, page in enumerate(pdf_reader.pages):
                    if i > 0:  # 添加页面分隔
                        markdown_content.append(f"\n---\n# 第 {i + 1} 页\n")

                    text = page.extract_text()
                    if text:
                        # 简单的段落处理
                        paragraphs = text.split("\n\n")
                        for paragraph in paragraphs:
                            paragraph = paragraph.strip()
                            if paragraph:
                                markdown_content.append(paragraph)
                                markdown_content.append("")

            return "\n".join(markdown_content)

        except Exception as e:
            raise Exception(f"PDF文件转换失败 (PyPDF2): {e}")

    @staticmethod
    def get_converter_for_extension(file_extension: str):
        """
        根据文件扩展名获取对应的转换器

        Args:
            file_extension: 文件扩展名（带点，如'.docx'）

        Returns:
            callable: 对应的转换函数，如果不需要转换则返回None
        """
        converters = {
            ".docx": FileConverter.docx_to_markdown,
            ".xlsx": FileConverter.xlsx_to_markdown,
            ".xls": FileConverter.xlsx_to_markdown,
            ".pdf": FileConverter.pdf_to_markdown,
        }

        return converters.get(file_extension.lower())
